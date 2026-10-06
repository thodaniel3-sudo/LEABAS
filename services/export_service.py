"""
LEBAS - Export Service

File:
    services/export_service.py

Purpose:
    Provides reusable functions for exporting LEBAS laboratory data,
    calculations, summaries, and reports to CSV and Excel files.

Design principles:
    - No experimental data is hard-coded.
    - Exports are generated from the current DataFrame/results.
    - Repeated sieve rows are preserved in raw-data exports.
    - Process-level summaries are deduplicated by sample_id + replicate.
    - Output folders are created automatically.
    - Existing files are not silently overwritten when a unique filename
      can be generated.
    - Excel files are formatted for practical laboratory/report use.

Supported exports:
    - Raw dataset CSV
    - Raw dataset Excel
    - Results CSV
    - Results Excel
    - Sieve analysis Excel
    - Beneficiation results Excel
    - Material balance Excel
    - Complete workbook containing multiple analysis sheets
"""


from __future__ import annotations

import logging
from pathlib import Path
from typing import Any, Dict, Iterable, Optional

import pandas as pd


logger = logging.getLogger(__name__)


# ============================================================================
# DEFAULT DIRECTORIES
# ============================================================================

PROJECT_ROOT = Path(
    __file__
).resolve().parent.parent

EXPORTS_DIR = PROJECT_ROOT / "exports"

REPORTS_DIR = PROJECT_ROOT / "reports"


# ============================================================================
# PROCESS-LEVEL COLUMNS
# ============================================================================

PROCESS_LEVEL_COLUMNS = [
    "feed_mass",
    "feed_grade",
    "feed_moisture",
    "crusher_type",
    "crush_time",
    "crush_power",
    "crush_feed_size",
    "crush_product_size",
    "grinding_time",
    "grinding_mill_type",
    "grinding_media_mass",
    "grinding_speed",
    "grinding_feed_size",
    "grinding_product_size",
    "beneficiation_method",
    "beneficiation_time",
    "reagent_type",
    "reagent_dosage",
    "ph_level",
    "pulp_density",
    "temperature",
    "concentrate_mass",
    "concentrate_grade",
    "tailings_mass",
    "tailings_grade",
    "losses_mass",
    "water_added",
    "solid_recovery",
]


# ============================================================================
# DIRECTORY HELPERS
# ============================================================================

def ensure_export_directory(
    directory: Optional[str | Path] = None,
) -> Path:
    """
    Create and return the export directory.
    """

    if directory is None:

        path = EXPORTS_DIR

    else:

        path = Path(
            directory
        ).expanduser().resolve()

    path.mkdir(
        parents=True,
        exist_ok=True,
    )

    return path


def ensure_reports_directory(
    directory: Optional[str | Path] = None,
) -> Path:
    """
    Create and return the reports directory.
    """

    if directory is None:

        path = REPORTS_DIR

    else:

        path = Path(
            directory
        ).expanduser().resolve()

    path.mkdir(
        parents=True,
        exist_ok=True,
    )

    return path


# ============================================================================
# FILENAME HELPERS
# ============================================================================

def _sanitize_filename(
    filename: str,
) -> str:
    """
    Remove characters that are unsafe in Windows filenames.
    """

    unsafe = '<>:"/\\|?*'

    cleaned = "".join(
        "_"
        if character in unsafe
        else character
        for character in str(filename)
    )

    cleaned = cleaned.strip()

    if not cleaned:

        cleaned = "export"

    return cleaned


def _ensure_extension(
    filename: str | Path,
    extension: str,
) -> str:
    """
    Ensure a filename has the requested extension.
    """

    extension = extension.lower()

    if not extension.startswith("."):

        extension = "." + extension

    name = Path(
        str(filename)
    ).name

    if not name.lower().endswith(
        extension
    ):

        name += extension

    return _sanitize_filename(
        name
    )


def _unique_path(
    path: Path,
) -> Path:
    """
    Return a unique path without overwriting an existing file.

    Example:

        results.xlsx
        results_1.xlsx
        results_2.xlsx
    """

    if not path.exists():

        return path

    counter = 1

    while True:

        candidate = path.with_name(
            f"{path.stem}_{counter}{path.suffix}"
        )

        if not candidate.exists():

            return candidate

        counter += 1


# ============================================================================
# DATAFRAME PREPARATION
# ============================================================================

def _prepare_dataframe(
    df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Validate and copy a DataFrame before export.

    The original DataFrame is never modified.
    """

    if df is None:

        raise ValueError(
            "Cannot export a None DataFrame."
        )

    if not isinstance(
        df,
        pd.DataFrame,
    ):

        raise TypeError(
            "Export data must be a pandas DataFrame."
        )

    if df.empty:

        raise ValueError(
            "Cannot export an empty DataFrame."
        )

    result = df.copy()

    # Normalize column labels without changing the user's data values.
    result.columns = [
        str(column).strip()
        for column in result.columns
    ]

    return result


def _prepare_process_dataframe(
    df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Reduce repeated sieve rows to one process-level record per
    sample_id + replicate.

    This is essential for:
        feed mass
        beneficiation masses
        crushing values
        grinding values
        etc.

    It prevents repeated sieve rows from being counted multiple times.
    """

    working = _prepare_dataframe(
        df
    )

    if not {
        "sample_id",
        "replicate",
    }.issubset(
        working.columns
    ):

        return working

    available = [
        column
        for column in PROCESS_LEVEL_COLUMNS
        if column in working.columns
    ]

    if available:

        working[
            "_process_completeness"
        ] = working[
            available
        ].notna().sum(axis=1)

        working = working.sort_values(
            "_process_completeness",
            ascending=False,
            kind="stable",
        )

        working = (
            working
            .groupby(
                [
                    "sample_id",
                    "replicate",
                ],
                as_index=False,
                sort=False,
            )
            .first()
        )

        working = working.drop(
            columns=[
                "_process_completeness"
            ],
            errors="ignore",
        )

    else:

        working = (
            working
            .groupby(
                [
                    "sample_id",
                    "replicate",
                ],
                as_index=False,
                sort=False,
            )
            .first()
        )

    return working.reset_index(
        drop=True
    )


# ============================================================================
# CSV EXPORT
# ============================================================================

def export_csv(
    df: pd.DataFrame,
    filename: str | Path = "lebas_export.csv",
    directory: Optional[str | Path] = None,
    index: bool = False,
    unique_filename: bool = True,
) -> Path:
    """
    Export a DataFrame to CSV.

    Returns:
        Path to the created CSV file.
    """

    data = _prepare_dataframe(
        df
    )

    output_dir = ensure_export_directory(
        directory
    )

    filename = _ensure_extension(
        filename,
        ".csv",
    )

    output_path = (
        output_dir
        / filename
    )

    if unique_filename:

        output_path = _unique_path(
            output_path
        )

    data.to_csv(
        output_path,
        index=index,
        encoding="utf-8-sig",
    )

    logger.info(
        "CSV export created: %s",
        output_path,
    )

    return output_path


# ============================================================================
# EXCEL FORMATTING
# ============================================================================

def _format_excel_writer(
    writer: pd.ExcelWriter,
    dataframe: pd.DataFrame,
    sheet_name: str,
) -> None:
    """
    Apply practical formatting to an Excel worksheet.

    Requires openpyxl.
    """

    try:

        worksheet = writer.sheets[
            sheet_name
        ]

        # Freeze the header row.
        worksheet.freeze_panes = "A2"

        # Enable filters.
        if len(dataframe.columns) > 0:

            last_column = (
                len(dataframe.columns)
            )

            last_row = (
                len(dataframe)
                + 1
            )

            from openpyxl.utils import get_column_letter

            last_column_letter = (
                get_column_letter(
                    last_column
                )
            )

            worksheet.auto_filter.ref = (
                f"A1:{last_column_letter}{last_row}"
            )

        # Header formatting.
        from openpyxl.styles import Font

        for cell in worksheet[1]:

            cell.font = Font(
                bold=True
            )

        # Reasonable column widths.
        for column_index, column in enumerate(
            dataframe.columns,
            start=1,
        ):

            values = [
                str(column)
            ]

            if not dataframe.empty:

                sample_values = (
                    dataframe[
                        column
                    ]
                    .head(100)
                    .fillna("")
                    .astype(str)
                    .tolist()
                )

                values.extend(
                    sample_values
                )

            width = max(
                len(value)
                for value in values
            )

            width = min(
                max(
                    width + 2,
                    10,
                ),
                40,
            )

            worksheet.column_dimensions[
                get_column_letter(
                    column_index
                )
            ].width = width

    except Exception as exc:

        logger.warning(
            "Excel formatting could not be fully applied: %s",
            exc,
        )


# ============================================================================
# EXCEL EXPORT
# ============================================================================

def export_excel(
    df: pd.DataFrame,
    filename: str | Path = "lebas_export.xlsx",
    directory: Optional[str | Path] = None,
    sheet_name: str = "Data",
    index: bool = False,
    unique_filename: bool = True,
) -> Path:
    """
    Export a DataFrame to an Excel workbook.
    """

    data = _prepare_dataframe(
        df
    )

    output_dir = ensure_export_directory(
        directory
    )

    filename = _ensure_extension(
        filename,
        ".xlsx",
    )

    output_path = (
        output_dir
        / filename
    )

    if unique_filename:

        output_path = _unique_path(
            output_path
        )

    try:

        with pd.ExcelWriter(
            output_path,
            engine="openpyxl",
        ) as writer:

            safe_sheet_name = (
                str(sheet_name)
                .strip()
                [:31]
                or "Data"
            )

            data.to_excel(
                writer,
                sheet_name=safe_sheet_name,
                index=index,
            )

            _format_excel_writer(
                writer,
                data,
                safe_sheet_name,
            )

    except ImportError as exc:

        raise ImportError(
            "Excel export requires openpyxl. "
            "Install it with: pip install openpyxl"
        ) from exc

    logger.info(
        "Excel export created: %s",
        output_path,
    )

    return output_path


# ============================================================================
# PROCESS RESULTS EXPORT
# ============================================================================

def export_process_results(
    df: pd.DataFrame,
    filename: str | Path = "lebas_process_results.xlsx",
    directory: Optional[str | Path] = None,
) -> Path:
    """
    Export process-level results.

    Repeated sieve rows are reduced to one row per sample/replicate.
    """

    process_df = _prepare_process_dataframe(
        df
    )

    return export_excel(
        process_df,
        filename=filename,
        directory=directory,
        sheet_name="Process Results",
    )


# ============================================================================
# SIEVE EXPORT
# ============================================================================

def export_sieve_results(
    df: pd.DataFrame,
    filename: str | Path = "lebas_sieve_analysis.xlsx",
    directory: Optional[str | Path] = None,
) -> Path:
    """
    Export sieve-analysis records.

    Unlike process-level exports, sieve rows are intentionally preserved.
    """

    data = _prepare_dataframe(
        df
    )

    sieve_columns = [
        column
        for column in [
            "sample_id",
            "material_type",
            "experiment_date",
            "operator",
            "replicate",
            "sieve_size",
            "mass_retained",
        ]
        if column in data.columns
    ]

    if not sieve_columns:

        raise ValueError(
            "The dataset does not contain enough sieve-analysis columns."
        )

    sieve_df = data[
        sieve_columns
    ].copy()

    # Calculate useful sieve-analysis values when possible.
    if {
        "sample_id",
        "replicate",
        "sieve_size",
        "mass_retained",
    }.issubset(
        sieve_df.columns
    ):

        sieve_df[
            "mass_retained"
        ] = pd.to_numeric(
            sieve_df[
                "mass_retained"
            ],
            errors="coerce",
        )

        group_columns = [
            "sample_id",
            "replicate",
        ]

        total_retained = (
            sieve_df
            .groupby(
                group_columns,
                dropna=False,
            )[
                "mass_retained"
            ]
            .transform(
                "sum"
            )
        )

        sieve_df[
            "percent_retained"
        ] = (
            sieve_df[
                "mass_retained"
            ]
            / total_retained
            * 100.0
        )

        sieve_df[
            "cumulative_retained"
        ] = (
            sieve_df
            .sort_values(
                [
                    "sample_id",
                    "replicate",
                    "sieve_size",
                ],
                ascending=[
                    True,
                    True,
                    False,
                ],
            )
            .groupby(
                group_columns,
                dropna=False,
            )[
                "mass_retained"
            ]
            .cumsum()
        )

        sieve_df[
            "cumulative_percent_retained"
        ] = (
            sieve_df[
                "cumulative_retained"
            ]
            / total_retained
            * 100.0
        )

        sieve_df[
            "percent_passing"
        ] = (
            100.0
            - sieve_df[
                "cumulative_percent_retained"
            ]
        )

    return export_excel(
        sieve_df,
        filename=filename,
        directory=directory,
        sheet_name="Sieve Analysis",
    )


# ============================================================================
# BENEFICIATION EXPORT
# ============================================================================

def export_beneficiation_results(
    df: pd.DataFrame,
    filename: str | Path = "lebas_beneficiation_results.xlsx",
    directory: Optional[str | Path] = None,
) -> Path:
    """
    Export beneficiation calculations.
    """

    process = _prepare_process_dataframe(
        df
    )

    numeric = [
        "feed_mass",
        "feed_grade",
        "concentrate_mass",
        "concentrate_grade",
        "tailings_mass",
        "tailings_grade",
        "losses_mass",
        "solid_recovery",
    ]

    for column in numeric:

        if column in process.columns:

            process[
                column
            ] = pd.to_numeric(
                process[
                    column
                ],
                errors="coerce",
            )

    if "feed_mass" in process.columns:

        feed = process[
            "feed_mass"
        ]

        concentrate = (
            process[
                "concentrate_mass"
            ]
            if "concentrate_mass"
            in process.columns
            else 0.0
        )

        process[
            "calculated_mass_recovery"
        ] = (
            concentrate
            / feed
            * 100.0
        )

        process[
            "calculated_yield"
        ] = (
            concentrate
            / feed
            * 100.0
        )

    if {
        "feed_mass",
        "feed_grade",
        "concentrate_mass",
        "concentrate_grade",
    }.issubset(
        process.columns
    ):

        denominator = (
            process[
                "feed_mass"
            ]
            * process[
                "feed_grade"
            ]
        )

        numerator = (
            process[
                "concentrate_mass"
            ]
            * process[
                "concentrate_grade"
            ]
        )

        process[
            "grade_recovery"
        ] = (
            numerator
            / denominator
            * 100.0
        )

        process[
            "enrichment_ratio"
        ] = (
            process[
                "concentrate_grade"
            ]
            / process[
                "feed_grade"
            ]
        )

    if "feed_mass" in process.columns:

        concentrate = (
            process[
                "concentrate_mass"
            ]
            if "concentrate_mass"
            in process.columns
            else 0.0
        )

        tailings = (
            process[
                "tailings_mass"
            ]
            if "tailings_mass"
            in process.columns
            else 0.0
        )

        losses = (
            process[
                "losses_mass"
            ]
            if "losses_mass"
            in process.columns
            else 0.0
        )

        process[
            "mass_balance"
        ] = (
            process[
                "feed_mass"
            ]
            - concentrate
            - tailings
            - losses
        )

        process[
            "balance_error_percent"
        ] = (
            process[
                "mass_balance"
            ].abs()
            / process[
                "feed_mass"
            ].abs()
            * 100.0
        )

    return export_excel(
        process,
        filename=filename,
        directory=directory,
        sheet_name="Beneficiation",
    )


# ============================================================================
# MATERIAL BALANCE EXPORT
# ============================================================================

def export_material_balance(
    df: pd.DataFrame,
    filename: str | Path = "lebas_material_balance.xlsx",
    directory: Optional[str | Path] = None,
) -> Path:
    """
    Export material-balance calculations.
    """

    process = _prepare_process_dataframe(
        df
    )

    numeric_columns = [
        "feed_mass",
        "concentrate_mass",
        "tailings_mass",
        "losses_mass",
    ]

    for column in numeric_columns:

        if column not in process.columns:

            process[
                column
            ] = 0.0

        else:

            process[
                column
            ] = pd.to_numeric(
                process[
                    column
                ],
                errors="coerce",
            )

    process[
        "output_mass"
    ] = (
        process[
            "concentrate_mass"
        ].fillna(0)
        + process[
            "tailings_mass"
        ].fillna(0)
        + process[
            "losses_mass"
        ].fillna(0)
    )

    process[
        "balance_difference"
    ] = (
        process[
            "feed_mass"
        ]
        - process[
            "output_mass"
        ]
    )

    process[
        "closure_percent"
    ] = (
        process[
            "output_mass"
        ]
        / process[
            "feed_mass"
        ]
        * 100.0
    )

    process[
        "mass_recovery_percent"
    ] = (
        process[
            "concentrate_mass"
        ]
        / process[
            "feed_mass"
        ]
        * 100.0
    )

    process[
        "balance_error_percent"
    ] = (
        process[
            "balance_difference"
        ].abs()
        / process[
            "feed_mass"
        ].abs()
        * 100.0
    )

    return export_excel(
        process,
        filename=filename,
        directory=directory,
        sheet_name="Material Balance",
    )


# ============================================================================
# COMPLETE WORKBOOK
# ============================================================================

def export_complete_workbook(
    df: pd.DataFrame,
    filename: str | Path = "LEBAS_complete_analysis.xlsx",
    directory: Optional[str | Path] = None,
    additional_sheets: Optional[
        Dict[str, pd.DataFrame]
    ] = None,
) -> Path:
    """
    Create one comprehensive Excel workbook containing multiple
    LEBAS analysis sheets.

    Sheets may include:
        Raw Data
        Process Results
        Sieve Analysis
        Beneficiation
        Material Balance
        Additional user-provided analysis sheets
    """

    data = _prepare_dataframe(
        df
    )

    process = _prepare_process_dataframe(
        data
    )

    output_dir = ensure_export_directory(
        directory
    )

    filename = _ensure_extension(
        filename,
        ".xlsx",
    )

    output_path = (
        output_dir
        / filename
    )

    output_path = _unique_path(
        output_path
    )

    try:

        with pd.ExcelWriter(
            output_path,
            engine="openpyxl",
        ) as writer:

            # --------------------------------------------------------------
            # RAW DATA
            # --------------------------------------------------------------

            data.to_excel(
                writer,
                sheet_name="Raw Data",
                index=False,
            )

            _format_excel_writer(
                writer,
                data,
                "Raw Data",
            )

            # --------------------------------------------------------------
            # PROCESS RESULTS
            # --------------------------------------------------------------

            process.to_excel(
                writer,
                sheet_name="Process Results",
                index=False,
            )

            _format_excel_writer(
                writer,
                process,
                "Process Results",
            )

            # --------------------------------------------------------------
            # SIEVE ANALYSIS
            # --------------------------------------------------------------

            try:

                sieve_data = _prepare_dataframe(
                    data
                )

                sieve_columns = [
                    column
                    for column in [
                        "sample_id",
                        "material_type",
                        "experiment_date",
                        "operator",
                        "replicate",
                        "sieve_size",
                        "mass_retained",
                    ]
                    if column in sieve_data.columns
                ]

                if sieve_columns:

                    sieve_data = sieve_data[
                        sieve_columns
                    ].copy()

                    sieve_data.to_excel(
                        writer,
                        sheet_name="Sieve Analysis",
                        index=False,
                    )

                    _format_excel_writer(
                        writer,
                        sieve_data,
                        "Sieve Analysis",
                    )

            except Exception as exc:

                logger.warning(
                    "Sieve sheet could not be generated: %s",
                    exc,
                )

            # --------------------------------------------------------------
            # BENEFICIATION
            # --------------------------------------------------------------

            try:

                beneficiation = (
                    export_beneficiation_dataframe(
                        data
                    )
                )

                beneficiation.to_excel(
                    writer,
                    sheet_name="Beneficiation",
                    index=False,
                )

                _format_excel_writer(
                    writer,
                    beneficiation,
                    "Beneficiation",
                )

            except Exception as exc:

                logger.warning(
                    "Beneficiation sheet could not be generated: %s",
                    exc,
                )

            # --------------------------------------------------------------
            # MATERIAL BALANCE
            # --------------------------------------------------------------

            try:

                balance = (
                    export_material_balance_dataframe(
                        data
                    )
                )

                balance.to_excel(
                    writer,
                    sheet_name="Material Balance",
                    index=False,
                )

                _format_excel_writer(
                    writer,
                    balance,
                    "Material Balance",
                )

            except Exception as exc:

                logger.warning(
                    "Material balance sheet could not be generated: %s",
                    exc,
                )

            # --------------------------------------------------------------
            # ADDITIONAL SHEETS
            # --------------------------------------------------------------

            if additional_sheets:

                for requested_name, sheet_df in (
                    additional_sheets.items()
                ):

                    if not isinstance(
                        sheet_df,
                        pd.DataFrame,
                    ):

                        logger.warning(
                            "Skipping additional sheet '%s': "
                            "value is not a DataFrame.",
                            requested_name,
                        )

                        continue

                    if sheet_df.empty:

                        logger.warning(
                            "Skipping empty additional sheet '%s'.",
                            requested_name,
                        )

                        continue

                    sheet_name = (
                        str(requested_name)
                        .strip()
                        [:31]
                        or "Analysis"
                    )

                    # Prevent duplicate sheet names.
                    if sheet_name in writer.sheets:

                        base_name = sheet_name[
                            :25
                        ]

                        counter = 1

                        while (
                            f"{base_name}_{counter}"
                            in writer.sheets
                        ):

                            counter += 1

                        sheet_name = (
                            f"{base_name}_{counter}"
                        )

                    additional = sheet_df.copy()

                    additional.to_excel(
                        writer,
                        sheet_name=sheet_name,
                        index=False,
                    )

                    _format_excel_writer(
                        writer,
                        additional,
                        sheet_name,
                    )

    except ImportError as exc:

        raise ImportError(
            "Excel export requires openpyxl. "
            "Install it with: pip install openpyxl"
        ) from exc

    logger.info(
        "Complete LEBAS workbook created: %s",
        output_path,
    )

    return output_path


# ============================================================================
# DATAFRAME-LEVEL ANALYSIS HELPERS
# ============================================================================

def export_beneficiation_dataframe(
    df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Prepare beneficiation results without writing a file.

    Used internally by the complete workbook exporter and also useful
    to other LEBAS modules.
    """

    process = _prepare_process_dataframe(
        df
    )

    numeric = [
        "feed_mass",
        "feed_grade",
        "concentrate_mass",
        "concentrate_grade",
        "tailings_mass",
        "tailings_grade",
        "losses_mass",
        "solid_recovery",
    ]

    for column in numeric:

        if column in process.columns:

            process[
                column
            ] = pd.to_numeric(
                process[
                    column
                ],
                errors="coerce",
            )

    if {
        "feed_mass",
        "concentrate_mass",
    }.issubset(
        process.columns
    ):

        process[
            "calculated_mass_recovery"
        ] = (
            process[
                "concentrate_mass"
            ]
            / process[
                "feed_mass"
            ]
            * 100.0
        )

        process[
            "calculated_yield"
        ] = process[
            "calculated_mass_recovery"
        ]

    if {
        "feed_mass",
        "feed_grade",
        "concentrate_mass",
        "concentrate_grade",
    }.issubset(
        process.columns
    ):

        denominator = (
            process[
                "feed_mass"
            ]
            * process[
                "feed_grade"
            ]
        )

        numerator = (
            process[
                "concentrate_mass"
            ]
            * process[
                "concentrate_grade"
            ]
        )

        process[
            "grade_recovery"
        ] = (
            numerator
            / denominator
            * 100.0
        )

        process[
            "enrichment_ratio"
        ] = (
            process[
                "concentrate_grade"
            ]
            / process[
                "feed_grade"
            ]
        )

    if {
        "feed_mass",
        "concentrate_mass",
    }.issubset(
        process.columns
    ):

        tailings = (
            process[
                "tailings_mass"
            ]
            if "tailings_mass"
            in process.columns
            else 0.0
        )

        losses = (
            process[
                "losses_mass"
            ]
            if "losses_mass"
            in process.columns
            else 0.0
        )

        process[
            "mass_balance"
        ] = (
            process[
                "feed_mass"
            ]
            - process[
                "concentrate_mass"
            ].fillna(0)
            - (
                tailings
                if isinstance(
                    tailings,
                    pd.Series,
                )
                else 0
            )
            - (
                losses
                if isinstance(
                    losses,
                    pd.Series,
                )
                else 0
            )
        )

        process[
            "balance_error_percent"
        ] = (
            process[
                "mass_balance"
            ].abs()
            / process[
                "feed_mass"
            ].abs()
            * 100.0
        )

    return process


def export_material_balance_dataframe(
    df: pd.DataFrame,
) -> pd.DataFrame:
    """
    Prepare material-balance results without writing a file.
    """

    process = _prepare_process_dataframe(
        df
    )

    numeric_columns = [
        "feed_mass",
        "concentrate_mass",
        "tailings_mass",
        "losses_mass",
    ]

    for column in numeric_columns:

        if column not in process.columns:

            process[
                column
            ] = 0.0

        else:

            process[
                column
            ] = pd.to_numeric(
                process[
                    column
                ],
                errors="coerce",
            )

    process[
        "output_mass"
    ] = (
        process[
            "concentrate_mass"
        ].fillna(0)
        + process[
            "tailings_mass"
        ].fillna(0)
        + process[
            "losses_mass"
        ].fillna(0)
    )

    process[
        "balance_difference"
    ] = (
        process[
            "feed_mass"
        ]
        - process[
            "output_mass"
        ]
    )

    process[
        "closure_percent"
    ] = (
        process[
            "output_mass"
        ]
        / process[
            "feed_mass"
        ]
        * 100.0
    )

    process[
        "mass_recovery_percent"
    ] = (
        process[
            "concentrate_mass"
        ]
        / process[
            "feed_mass"
        ]
        * 100.0
    )

    process[
        "balance_error_percent"
    ] = (
        process[
            "balance_difference"
        ].abs()
        / process[
            "feed_mass"
        ].abs()
        * 100.0
    )

    return process


# ============================================================================
# TEXT EXPORT
# ============================================================================

def export_text(
    text: str,
    filename: str | Path = "lebas_report.txt",
    directory: Optional[str | Path] = None,
    unique_filename: bool = True,
) -> Path:
    """
    Save plain text content to the reports directory.
    """

    output_dir = ensure_reports_directory(
        directory
    )

    filename = _ensure_extension(
        filename,
        ".txt",
    )

    output_path = (
        output_dir
        / filename
    )

    if unique_filename:

        output_path = _unique_path(
            output_path
        )

    output_path.write_text(
        str(text),
        encoding="utf-8",
    )

    logger.info(
        "Text export created: %s",
        output_path,
    )

    return output_path


# ============================================================================
# EXPORT DIRECTORY LISTING
# ============================================================================

def list_exports(
    directory: Optional[str | Path] = None,
) -> list[Path]:
    """
    Return files currently present in the export directory.
    """

    output_dir = ensure_export_directory(
        directory
    )

    return sorted(
        [
            path
            for path in output_dir.iterdir()
            if path.is_file()
        ],
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )


# ============================================================================
# DELETE EXPORT
# ============================================================================

def delete_export(
    filepath: str | Path,
) -> bool:
    """
    Delete a generated export file.

    Returns:
        True if deleted successfully.
        False if the file did not exist.

    Raises:
        OSError if deletion fails.
    """

    path = Path(
        filepath
    ).expanduser().resolve()

    if not path.exists():

        return False

    if not path.is_file():

        raise ValueError(
            f"Export path is not a file: {path}"
        )

    path.unlink()

    logger.info(
        "Export deleted: %s",
        path,
    )

    return True


# ============================================================================
# BACKWARD-COMPATIBILITY ALIASES
# ============================================================================

save_csv = export_csv

save_excel = export_excel

export_to_csv = export_csv

export_to_excel = export_excel

create_excel_export = export_excel

create_csv_export = export_csv


# ============================================================================
# PUBLIC API
# ============================================================================

__all__ = [
    "PROJECT_ROOT",
    "EXPORTS_DIR",
    "REPORTS_DIR",
    "ensure_export_directory",
    "ensure_reports_directory",
    "export_csv",
    "export_excel",
    "export_process_results",
    "export_sieve_results",
    "export_beneficiation_results",
    "export_material_balance",
    "export_complete_workbook",
    "export_beneficiation_dataframe",
    "export_material_balance_dataframe",
    "export_text",
    "list_exports",
    "delete_export",
    "save_csv",
    "save_excel",
    "export_to_csv",
    "export_to_excel",
    "create_excel_export",
    "create_csv_export",
]